from io import BytesIO
from decimal import Decimal, ROUND_HALF_UP
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pages import documents

def euro(value):
    amount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP,)
    return f"{amount:,.2f}".replace(",", " ").replace(".", ",")

def cellBackground(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    back = OxmlElement("w:shd")
    back.set(qn("w:fill"), color)
    properties.append(back)

def addLign(tab, values, bold=False, color=None):
    cells = tab.add_row().cells
    for index, value in enumerate(values):
        cell = cells[index]
        cell.text = str(value)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        if color:
            cellBackground(cell, color)

        for paragraph in cell.paragraphs:
            if index >= 2:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)
                    run.bold = bold
    return cells

def createSell(articles, project="", offerDate="", validationDate="",):
    columns = {
            "Lot",
            "Nom du lot",
            "Bordereau",
            "Désignation",
            "Total vente HT",
    }

    missing = columns - set(articles.columns)
    if missing:
        raise ValueError("Colonnes manquantes : " + ", ".join(sorted(missing)))

    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.paege_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    titleText = title.add_run("DÉCOMPOSITION DU PRIX")
    titleText.bold = True;
    titleText.font.size = Pt(14)
    titleText.font.color.rgb = RGBColor(0,0,0)
    
    doc.add_paragraph("")

    for lib, val in [
        ("Nom du projet", project),
        ("Date de remise de l'offre", offerDate),
        ("Date de validité de l'offre", validationDate),
        ("Taux devise", "1 Euro = 1 Euro"),
    ]:
        lign = doc.add_paragraph()
        lign.paragraph_format.space_after = Pt(2)
        lign.add_run(f"{lib} : ").bold = True
        lign.add_run(str(val))

    doc.add_paragraph("")
    tab = doc.add_table(rows=0, cols=4)
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False

    spaces = [Cm(1.4), Cm(11.0), Cm(3.0), Cm(3.0)]

    for column, space in zip(tab.columns, spaces):
        column.width = space

    addLign(tab, ["N°", "Désignation", "Quantité", "Total Euro H.T."],
            bold=True, color="E8EDF2",)

    generalTotal = Decimal("0.00")

    for (num, name), lot in articles.groupby(
            ["Lot", "Nom du lot"],
            sort=False
    ):
        addLign(tab, [num, name, "", ""],bold=True, color="D9E7F4",)

        totalLot = Decimal("0.00")

        for nameBorder, border in lot.groupby("Bordereau", sort=False,):
            if nameBorder != "Sans bordereau":
                addLign(tab, ["", nameBorder, "",""], bold=True, color="F0F3F6",)
            for _, article in border.iterrows():
                amount = Decimal(str(article["Total vente HT"])
                    ).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP,)
                totalLot += amount
                quantity = article.get("Quantité", "")
                if quantity is None:
                    quantity = ""

                addLign(tab, [
                    article.get("Code", ""),
                    article["Désignation"],
                    quantity, euro(amount),
                    ],)
        generalTotal += totalLot
    addLign(tab,
        ["", "MONTANT DU CONTRAT HT", "", euro(generalTotal)],
        bold=True,
        color="D9E7F4"
    )

    doc.add_paragraph("")
    footer = doc.add_paragraph("Montants exprimés en euros hors taxes.")
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    output = BytesIO()
    doc.save(output)
    return output.getvalue()

